import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from datetime import datetime
from copy import copy
import shutil

src = 'whiteshirt_interview_master_SANITIZED.xlsx'
dst = 'outputs/whiteshirt_interview_master_WORKING.xlsx'

# Always start from fresh copy to keep fixes idempotent
shutil.copy(src, dst)

wb = openpyxl.load_workbook(dst)
TS = datetime.utcnow().strftime('%Y-%m-%d %H:%M:%SZ')
YELLOW = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
log = []

def col_idx(ws, name, header_row=4):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v and str(v).strip() == name:
            return c
    raise ValueError(f"Column not found: {name} in {ws.title}")

def set_cell(ws, row, col, new_val, fix_id, issue, evidence, reason):
    old_val = ws.cell(row=row, column=col).value
    if old_val == new_val:
        return False
    ws.cell(row=row, column=col).value = new_val
    ws.cell(row=row, column=col).fill = YELLOW
    log.append({
        'Fix_ID': fix_id, 'Timestamp': TS, 'Issue': issue,
        'Sheet': ws.title, 'Row': row,
        'Column': get_column_letter(col),
        'Column_Name': str(ws.cell(row=4, column=col).value or ''),
        'Old_Value': str(old_val)[:200] if old_val is not None else '',
        'New_Value': str(new_val)[:200] if new_val is not None else '',
        'Evidence': evidence, 'Reason': reason,
    })
    return True

def log_info(fix_id, issue, evidence, reason):
    log.append({
        'Fix_ID': fix_id, 'Timestamp': TS, 'Issue': issue,
        'Sheet': '(info only)', 'Row': '', 'Column': '', 'Column_Name': '',
        'Old_Value': '', 'New_Value': '',
        'Evidence': evidence, 'Reason': reason,
    })

# ==========================================
# P1: Insert missing DAILY_ACCOUNT row
# ==========================================
print("== P1 ==")
ws = wb['DAILY_ACCOUNT']
insert_at = 31  # between Slot 6 (row 30) and Slot 8 (row 31)

# snapshot formatting from row below
fmt = {c: copy(ws.cell(row=insert_at, column=c)._style) for c in range(1, ws.max_column + 1)}
ws.insert_rows(insert_at, 1)
for c in range(1, ws.max_column + 1):
    ws.cell(row=insert_at, column=c)._style = fmt[c]

fill_map = {
    'Date': datetime(2026, 7, 23),
    'Capture Time': '20:30',
    'Comment Account': 'Account Slot 7',
    'Account Group': 'B',
    'Account Action Status': 'Active',
    'Test Group': 'Advanced',
    'Data Owner': 'Interview Candidate',
    'QC Status': 'Review',
    'Notes': 'Reconstructed from import ' + TS + ' - verify remaining metrics',
    'Import Batch ID': 'BATCH-ACC-D3',
    'Reconciliation Status': 'Reconstructed',
}
for name, val in fill_map.items():
    c = col_idx(ws, name)
    ws.cell(row=insert_at, column=c).value = val
    ws.cell(row=insert_at, column=c).fill = YELLOW

log.append({
    'Fix_ID': 'P1', 'Timestamp': TS, 'Issue': 'Missing DAILY_ACCOUNT row',
    'Sheet': 'DAILY_ACCOUNT', 'Row': insert_at, 'Column': '(inserted)',
    'Column_Name': '(new row)', 'Old_Value': 'missing',
    'New_Value': '2026-07-23 / Account Slot 7',
    'Evidence': 'Row count 39 vs expected 40',
    'Reason': 'Completeness - 10 slots per day required',
})
print(f"  inserted row at Excel row {insert_at}")

# ==========================================
# P2: IC-0012 collision on Slot 3
# ==========================================
print("== P2 ==")
ws = wb['COMMENT_EXECUTION']
cid = col_idx(ws, 'Comment ID')
acc = col_idx(ws, 'Comment Account')
qc  = col_idx(ws, 'QC Exception Flag')

# Slot 3 at Excel row 17 → rename to IC-0013
set_cell(ws, 17, cid, 'IC-0013', 'P2',
         'Comment ID collision',
         'IC-0012 duplicated; IC-0013 missing in 1..100',
         'Row belongs to Account Slot 3')

# Reset false-positive QC flag for Slot 3 (was triggered by ID collision)
set_cell(ws, 17, qc, 'OK', 'P2b',
         'False positive QC flag',
         'Slot 3 text is normal (image-context fit)',
         'Flag likely triggered by ID collision with direct-sell row')

# Propagate to COMMENT_RECHECK
ws = wb['COMMENT_RECHECK']
rcid = col_idx(ws, 'Comment ID')
racc = col_idx(ws, 'Comment Account')
n_re = 0
for row in range(5, ws.max_row + 1):
    if ws.cell(row=row, column=rcid).value == 'IC-0012' and \
       ws.cell(row=row, column=racc).value == 'Account Slot 3':
        set_cell(ws, row, rcid, 'IC-0013', 'P2c',
                 'ID collision propagated to recheck',
                 f'Row {row} matches Account Slot 3 IC-0012',
                 'Align with corrected COMMENT_EXECUTION ID')
        n_re += 1
print(f"  renamed 1 CE row + {n_re} recheck rows")

# ==========================================
# P3: Test Group sync
# ==========================================
print("== P3 ==")
ws_da = wb['DAILY_ACCOUNT']
da_d = col_idx(ws_da, 'Date')
da_a = col_idx(ws_da, 'Comment Account')
da_t = col_idx(ws_da, 'Test Group')
tg = {}
for r in range(5, ws_da.max_row + 1):
    d = ws_da.cell(row=r, column=da_d).value
    a = ws_da.cell(row=r, column=da_a).value
    t = ws_da.cell(row=r, column=da_t).value
    if d and a:
        tg[(str(d)[:10], str(a))] = t

ws = wb['COMMENT_EXECUTION']
cd = col_idx(ws, 'Date')
ca = col_idx(ws, 'Comment Account')
ct = col_idx(ws, 'Test Group')
n = 0
for row in range(5, ws.max_row + 1):
    d = str(ws.cell(row=row, column=cd).value)[:10]
    a = str(ws.cell(row=row, column=ca).value)
    exp = tg.get((d, a))
    cur = ws.cell(row=row, column=ct).value
    if exp and cur != exp:
        set_cell(ws, row, ct, exp, 'P3',
                 'Test Group mismatch vs DAILY_ACCOUNT',
                 f'DAILY_ACCOUNT({d},{a}) = {exp}',
                 'Sync to account-level source of truth')
        n += 1
print(f"  synced {n} Test Group values")

# ==========================================
# P4: Flag direct-sell rows not caught
# ==========================================
print("== P4 ==")
ws = wb['COMMENT_EXECUTION']
cid = col_idx(ws, 'Comment ID')
qc  = col_idx(ws, 'QC Exception Flag')
targets = {'IC-0027','IC-0071','IC-0088'}
n = 0
for row in range(5, ws.max_row + 1):
    if ws.cell(row=row, column=cid).value in targets:
        if set_cell(ws, row, qc, 'Review', 'P4',
                    'Direct-sell keyword not flagged',
                    'Text contains forbidden sell keyword',
                    'Policy: no direct selling per brief'):
            n += 1
print(f"  flagged {n} direct-sell rows")

# ==========================================
# P8: Reply Owner Type = 0 → blank
# ==========================================
print("== P8 ==")
ws = wb['COMMENT_EXECUTION']
ro = col_idx(ws, 'Reply Owner Type')
cid = col_idx(ws, 'Comment ID')
n = 0
for row in range(5, ws.max_row + 1):
    if ws.cell(row=row, column=cid).value == 'IC-0044' and \
       ws.cell(row=row, column=ro).value == 0:
        set_cell(ws, row, ro, None, 'P8',
                 'Invalid Reply Owner Type value',
                 'No reply present; value was 0',
                 '0 not in enum {User, Creator, Mixed, None}')
        n += 1
print(f"  fixed {n} row")

# ==========================================
# P10: COMPLETENESS_CHECK status inconsistency
# ==========================================
print("== P10 ==")
ws = wb['COMPLETENESS_CHECK']
date_c = col_idx(ws, 'Date')
req_c  = col_idx(ws, 'Required Rows')
act_c  = col_idx(ws, 'Actual Rows')
st_c   = col_idx(ws, 'Daily Account Data Status')
n = 0
for row in range(5, ws.max_row + 1):
    act = ws.cell(row=row, column=act_c).value
    req = ws.cell(row=row, column=req_c).value
    st  = ws.cell(row=row, column=st_c).value
    if isinstance(act,(int,float)) and isinstance(req,(int,float)) and act < req and st != 'Review':
        set_cell(ws, row, st_c, 'Review', 'P10',
                 'Status inconsistent with row count',
                 f'Actual {act} < Required {req}',
                 'Force Review when rows missing')
        n += 1
print(f"  fixed {n} status")

# ==========================================
# P11: EVIDENCE_INDEX row 45→44
# ==========================================
print("== P11 ==")
ws = wb['EVIDENCE_INDEX']
eid = col_idx(ws, 'Evidence ID')
lr  = col_idx(ws, 'Linked Row')
for row in range(5, ws.max_row + 1):
    if ws.cell(row=row, column=eid).value == 'EVD-SYN-0040-POST':
        set_cell(ws, row, lr, 44, 'P11',
                 'Linked Row off by one',
                 'IC-0040 maps to Excel row 44 (= 40 + 4)',
                 'Off-by-one defect in evidence pointer')
        print(f"  fixed EVD-SYN-0040-POST at row {row}")
        break

# ==========================================
# P12: DAILY_BRAND Product Clicks mismatch
# ==========================================
print("== P12 ==")
ws = wb['DAILY_BRAND']
d_c = col_idx(ws, 'Date')
pc  = col_idx(ws, 'Product Clicks')
for row in range(5, ws.max_row + 1):
    d = ws.cell(row=row, column=d_c).value
    if d and str(d)[:10] == '2026-07-24':
        set_cell(ws, row, pc, 237, 'P12',
                 'Mismatch with input source',
                 'input brand_snapshot shows 237; master shows 242',
                 'Reconcile to input source of truth')
        print(f"  fixed Product Clicks 242 → 237 at row {row}")
        break

# ==========================================
# Informational notes (insufficient evidence)
# ==========================================
log_info('N1', 'Image Fit To Context = NaN 83 rows',
         '83/100 rows blank; only 17 marked Fit',
         'Cannot infer fit without manual review; needs verification')

log_info('N2', 'Cart Status = NaN 100 rows',
         'No source field in execution sheet',
         'Cart data source unknown; leave blank + verify upstream')

log_info('N3', 'Visibility pattern Visible→Missing 18 cases',
         'Suggest shadowban signal but no cause evidence',
         'Add fields First_Visible / Last_Visible / Change_Flag')

log_info('N4', 'IC-0031 flagged Review but text looks normal',
         'Text passes keyword check; reason not documented',
         'Keep flag; add QC_Reason field for traceability')

log_info('N5', 'Slot 1 + Slot 6 visibility 0%',
         '8 slot-days across 4 days',
         'Account health rule needed: pause if 0% for >= 2 days')

# ==========================================
# Add FIX_LOG sheet
# ==========================================
if 'FIX_LOG' in wb.sheetnames:
    del wb['FIX_LOG']
ws_log = wb.create_sheet('FIX_LOG')
headers = ['Fix_ID','Timestamp','Issue','Sheet','Row','Column','Column_Name',
           'Old_Value','New_Value','Evidence','Reason']
ws_log.append(headers)
for h in range(1, len(headers) + 1):
    cell = ws_log.cell(row=1, column=h)
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    ws_log.column_dimensions[get_column_letter(h)].width = 22
for entry in log:
    ws_log.append([entry.get(k, '') for k in headers])

wb.save(dst)
print()
print("=" * 60)
print(f"Fixes logged: {len(log)}")
print(f"Saved: {dst}")
print("=" * 60)
