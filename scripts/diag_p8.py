import openpyxl

wb = openpyxl.load_workbook('outputs/whiteshirt_interview_master_WORKING.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('outputs/whiteshirt_interview_master_WORKING.xlsx', data_only=True)
ws = wb['COMMENT_RECHECK']
ws_v = wb_v['COMMENT_RECHECK']

# Header row 4
cols = {str(ws.cell(row=4, column=c).value).strip(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=4, column=c).value}
print("COMMENT_RECHECK columns:")
for k, v in cols.items():
    print(f"  col {v:>2}: {k}")
print()

# Find IC-0044 rows
print("=" * 70)
print("Rows with Comment ID = IC-0044")
print("=" * 70)
cid_col = cols['Comment ID']
tp_col  = cols['Verification Timepoint']
own_col = cols['Lookup Reply Owner']
key_col = cols['Lookup Key']

for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=cid_col).value == 'IC-0044':
        tp = ws.cell(row=r, column=tp_col).value
        key_f = ws.cell(row=r, column=key_col).value
        own_f = ws.cell(row=r, column=own_col).value
        own_v = ws_v.cell(row=r, column=own_col).value
        print(f"row {r}: TP={tp!r}")
        print(f"  Lookup Key      formula: {key_f!r}")
        print(f"  Lookup Owner    formula: {own_f!r}")
        print(f"  Lookup Owner    cached : {own_v!r}")
        print()

# Search entire Lookup Reply Owner column for 0
print("=" * 70)
print("Any cell with Lookup Reply Owner == 0 (formula or literal)?")
print("=" * 70)
n = 0
for r in range(5, ws.max_row + 1):
    own_f = ws.cell(row=r, column=own_col).value
    own_v = ws_v.cell(row=r, column=own_col).value
    if own_f == 0 or own_f == '0' or own_v == 0 or own_v == '0':
        cid = ws.cell(row=r, column=cid_col).value
        tp = ws.cell(row=r, column=tp_col).value
        print(f"  row {r}: CID={cid} TP={tp} | formula={own_f!r} | cached={own_v!r}")
        n += 1
print(f"  total: {n} cells")
