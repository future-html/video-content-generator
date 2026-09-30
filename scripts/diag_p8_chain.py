import openpyxl

wb = openpyxl.load_workbook('outputs/whiteshirt_interview_master_WORKING.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('outputs/whiteshirt_interview_master_WORKING.xlsx', data_only=True)

print("=" * 70)
print("COMMENT_RECHECK rows 91, 92 — full chain")
print("=" * 70)
ws = wb['COMMENT_RECHECK']
ws_v = wb_v['COMMENT_RECHECK']

for r in [91, 92]:
    print(f"\n--- Row {r} ---")
    for c in [4, 5, 9, 14, 15, 16, 17, 18, 19, 20, 22, 23, 24, 25]:
        h = ws.cell(row=4, column=c).value
        v_f = ws.cell(row=r, column=c).value
        v_v = ws_v.cell(row=r, column=c).value
        flag = ""
        if v_v == 0 or v_v == '0':
            flag = "  ← ⚠ ZERO!"
        print(f"  col {c:>2} [{h}]: formula={v_f!r} | cached={v_v!r}{flag}")

print()
print("=" * 70)
print("Search for literal 0 anywhere in COMMENT_RECHECK (col R=18)")
print("=" * 70)
n = 0
for r in range(5, ws.max_row + 1):
    v_f = ws.cell(row=r, column=18).value
    v_v = ws_v.cell(row=r, column=18).value
    if v_f == 0 or v_f == '0' or v_v == 0 or v_v == '0':
        cid = ws.cell(row=r, column=4).value
        tp = ws.cell(row=r, column=19).value
        print(f"  row {r}: CID={cid} TP={tp} | formula={v_f!r} | cached={v_v!r}")
        n += 1
print(f"  total: {n}")

print()
print("=" * 70)
print("Search for literal 0 anywhere in COMMENT_EXECUTION (col AS=45)")
print("=" * 70)
ws2 = wb['COMMENT_EXECUTION']
ws2_v = wb_v['COMMENT_EXECUTION']
n = 0
for r in range(5, ws2.max_row + 1):
    v_f = ws2.cell(row=r, column=45).value
    v_v = ws2_v.cell(row=r, column=45).value
    if v_f == 0 or v_f == '0' or v_v == 0 or v_v == '0':
        cid = ws2.cell(row=r, column=2).value
        print(f"  row {r}: CID={cid} | formula={v_f!r} | cached={v_v!r}")
        n += 1
print(f"  total: {n}")
