import openpyxl
import json
from datetime import datetime
from pathlib import Path

SRC = 'outputs/whiteshirt_interview_master_WORKING.xlsx'
OUT = 'outputs/data.json'

def cell_val(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)):
        return v
    s = str(v)
    if s.startswith('='):
        return s  # keep formula
    return s

def sheet_to_records(ws, header_row=4):
    headers = []
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=c).value
        headers.append(str(h).strip() if h else f'col{c}')
    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {}
        empty = True
        for c in range(1, ws.max_column + 1):
            v = cell_val(ws.cell(row=r, column=c).value)
            row[headers[c-1]] = v
            if v not in (None, ''):
                empty = False
        if not empty:
            row['_row'] = r
            records.append(row)
    return headers, records

wb = openpyxl.load_workbook(SRC, data_only=False)
out = {'generated_at': datetime.utcnow().isoformat() + 'Z', 'sheets': {}}

for name in wb.sheetnames:
    ws = wb[name]
    if name == 'README_START_HERE':
        continue
    headers, records = sheet_to_records(ws, header_row=4)
    out['sheets'][name] = {
        'headers': headers,
        'row_count': len(records),
        'records': records,
    }
    print(f"  {name}: {len(records)} records, {len(headers)} columns")

Path(OUT).write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
print(f"\nWrote {OUT}")
print(f"Size: {Path(OUT).stat().st_size / 1024:.1f} KB")
