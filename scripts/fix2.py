import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone

dst = 'outputs/whiteshirt_interview_master_WORKING.xlsx'
wb = openpyxl.load_workbook(dst)
YELLOW = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
TS = datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%SZ')

def col_idx(ws, name, header_row=4):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v and str(v).strip() == name:
            return c
    raise ValueError(f"{name} not found in {ws.title}")

log = []
def add_log(fix_id, issue, sheet, row, col, cname, old, new, ev, reason):
    log.append([fix_id, TS, issue, sheet, row, col, cname,
                str(old)[:200] if old is not None else '',
                str(new)[:200] if new is not None else '',
                ev, reason])

# ============================================
# STEP 1: Restore + upgrade QC Exception Flag formula
# ============================================
print("="*70)
print("STEP 1: Restore QC Exception Flag with keyword check")
print("="*70)
ws = wb['COMMENT_EXECUTION']
qc = col_idx(ws, 'QC Exception Flag')

tpl = ('=IF(OR('
       'COUNTIF($B$5:$B$104,B{r})>1,'
       'AND(T{r}="Posted",BC{r}<>"Yes"),'
       'AND(OR(BJ{r}="Verified User Reply",BJ{r}="Mixed Reply"),Y{r}=""),'
       'ISNUMBER(SEARCH("ซื้อเลย",O{r})),'
       'ISNUMBER(SEARCH("โปรแรง",O{r})),'
       'ISNUMBER(SEARCH("ทักด่วน",O{r})),'
       'ISNUMBER(SEARCH("ทักแชต",O{r})),'
       'ISNUMBER(SEARCH("ลดราคา",O{r})),'
       'ISNUMBER(SEARCH("รีบซื้อ",O{r})),'
       'ISNUMBER(SEARCH("ส่งฟรี",O{r})),'
       'ISNUMBER(SEARCH("คลิกลิงก์",O{r}))'
       '),"Review","OK")')

n = 0
for r in range(5, ws.max_row + 1):
    cid = ws.cell(row=r, column=2).value
    if not cid or not str(cid).startswith('IC-'):
        continue
    cell = ws.cell(row=r, column=qc)
    old = cell.value
    new = tpl.format(r=r)
    if old != new:
        cell.value = new
        cell.fill = YELLOW
        n += 1
        add_log('P4fix', 'Restore formula + add keyword check',
                'COMMENT_EXECUTION', r, get_column_letter(qc),
                'QC Exception Flag', old, new,
                'Formula was overwritten; keywords now auto-flag',
                'Preserve formula integrity; flag direct-sell via SEARCH')
print(f"  formulas restored/modified: {n}")

# ============================================
# STEP 2: Diagnose + fix P2 recheck propagation
# ============================================
print()
print("="*70)
print("STEP 2: P2 - Recheck propagation")
print("="*70)
ws = wb['COMMENT_RECHECK']
rcid = col_idx(ws, 'Comment ID')
racc = col_idx(ws, 'Comment Account')
rtime = col_idx(ws, 'Verification Timepoint')

print("  IC-0012/IC-0013 rows currently in COMMENT_RECHECK:")
target_rows = []
for r in range(5, ws.max_row + 1):
    cid = ws.cell(row=r, column=rcid).value
    if cid in ('IC-0012', 'IC-0013'):
        acc = ws.cell(row=r, column=racc).value
        tp = ws.cell(row=r, column=rtime).value
        print(f"    row {r}: ID={cid} | Account={acc} | TP={tp}")
        if cid == 'IC-0012' and acc == 'Account Slot 3':
            target_rows.append(r)

if target_rows:
    for r in target_rows:
        old = ws.cell(row=r, column=rcid).value
        ws.cell(row=r, column=rcid).value = 'IC-0013'
        ws.cell(row=r, column=rcid).fill = YELLOW
        add_log('P2', 'Propagate ID rename to recheck',
                'COMMENT_RECHECK', r, get_column_letter(rcid),
                'Comment ID', old, 'IC-0013',
                'Matches Slot 3', 'Align with corrected CE ID')
    print(f"  recheck rows renamed: {len(target_rows)}")
else:
    print("  ⚠ No IC-0012+Slot3 rows found - check ID pattern")

# ============================================
# STEP 3: Diagnose + fix P8 (IC-0044 Reply Owner = 0)
# ============================================
print()
print("="*70)
print("STEP 3: P8 - IC-0044 Reply Owner Type")
print("="*70)
ws = wb['COMMENT_EXECUTION']
ro = col_idx(ws, 'Reply Owner Type')

fixed_p8 = 0
for r in range(5, ws.max_row + 1):
    if ws.cell(row=r, column=2).value == 'IC-0044':
        val = ws.cell(row=r, column=ro).value
        print(f"  row {r}: value={val!r} type={type(val).__name__}")
        # match both int 0 and str "0"
        if val == 0 or val == '0' or (isinstance(val, str) and val.strip() == '0'):
            ws.cell(row=r, column=ro).value = None
            ws.cell(row=r, column=ro).fill = YELLOW
            add_log('P8', 'Invalid Reply Owner Type value',
                    'COMMENT_EXECUTION', r, get_column_letter(ro),
                    'Reply Owner Type', val, '',
                    'Value 0 not in enum {User, Creator, Mixed, None}',
                    'No reply present; blank is correct')
            fixed_p8 += 1
print(f"  P8 fixed: {fixed_p8} row(s)")

# ============================================
# STEP 4: Append to FIX_LOG
# ============================================
print()
print("="*70)
print("STEP 4: Append corrective log to FIX_LOG")
print("="*70)
if 'FIX_LOG' in wb.sheetnames:
    ws_log = wb['FIX_LOG']
    for entry in log:
        ws_log.append(entry)
    print(f"  {len(log)} entries appended to FIX_LOG")

wb.save(dst)
print()
print("="*70)
print(f"Saved: {dst}")
print(f"Total corrective log entries: {len(log)}")
print("="*70)
