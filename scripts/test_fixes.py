"""Whiteshirt Fix Verification Suite v3 — formula-aware"""
import openpyxl
import sys

ORIG = 'whiteshirt_interview_master_SANITIZED.xlsx'
WORK = 'outputs/whiteshirt_interview_master_WORKING.xlsx'

PASS, FAIL = 0, 0
def test(name, cond, detail=''):
    global PASS, FAIL
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  ({detail})" if detail else ''))
    if cond: PASS += 1
    else:    FAIL += 1

def hdr(wb, sheet):
    ws = wb[sheet]
    return {str(ws.cell(row=4, column=c).value).strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=4, column=c).value}

def find_row(ws, col, val, start=5):
    for r in range(start, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == val:
            return r
    return None

def is_formula(v):
    return isinstance(v, str) and v.startswith('=')

print("=" * 70)
print("WHITESHIRT FIX VERIFICATION SUITE v3")
print("=" * 70)

wb = openpyxl.load_workbook(WORK, data_only=False)
wb_v = openpyxl.load_workbook(WORK, data_only=True)
wb_orig = openpyxl.load_workbook(ORIG, data_only=False)

# P1
print("\n[P1] DAILY_ACCOUNT")
ws = wb['DAILY_ACCOUNT']; h = hdr(wb, 'DAILY_ACCOUNT')
n = ws.max_row - 4
test("40 rows", n == 40, f"got {n}")
slots_23 = [ws.cell(row=r, column=h['Comment Account']).value
            for r in range(5, ws.max_row + 1)
            if str(ws.cell(row=r, column=h['Date']).value)[:10] == '2026-07-23']
test("2026-07-23 = 10 slots", len(slots_23) == 10, f"got {len(slots_23)}")
test("Slot 7 present", 'Account Slot 7' in slots_23)

# P2
print("\n[P2] Comment ID uniqueness")
ws = wb['COMMENT_EXECUTION']; h = hdr(wb, 'COMMENT_EXECUTION')
ids = [ws.cell(row=r, column=h['Comment ID']).value for r in range(5, ws.max_row + 1)]
test("100 unique", len(set(ids)) == 100, f"got {len(set(ids))}")
test("IC-0013 exists", 'IC-0013' in ids)
test("IC-0012 count = 1", ids.count('IC-0012') == 1)

# P3
print("\n[P3] Test Group sync")
ws_da = wb['DAILY_ACCOUNT']; h_da = hdr(wb, 'DAILY_ACCOUNT')
da_map = {(str(ws_da.cell(row=r, column=h_da['Date']).value)[:10],
           ws_da.cell(row=r, column=h_da['Comment Account']).value):
          ws_da.cell(row=r, column=h_da['Test Group']).value
          for r in range(5, ws_da.max_row + 1)}
mm = sum(1 for r in range(5, ws.max_row + 1)
         if (k := (str(ws.cell(row=r, column=h['Date']).value)[:10],
                   ws.cell(row=r, column=h['Comment Account']).value))
         and da_map.get(k) and ws.cell(row=r, column=h['Test Group']).value != da_map[k])
test("0 mismatch", mm == 0, f"got {mm}")

# P4
print("\n[P4] Direct-sell QC (formula-aware)")
for cid in ['IC-0027', 'IC-0071', 'IC-0088']:
    r = find_row(ws, h['Comment ID'], cid)
    if r:
        v = ws.cell(row=r, column=h['QC Exception Flag']).value
        if is_formula(v):
            has_kw = all(k in v for k in ['ซื้อเลย', 'ทักด่วน', 'ทักแชต', 'ลดราคา'])
            test(f"{cid} QC has keywords", has_kw)
        else:
            test(f"{cid} QC = Review", v == 'Review', f"got {v!r}")
qc_vals = [ws.cell(row=r, column=h['QC Exception Flag']).value for r in range(5, ws.max_row + 1)]
n_f = sum(1 for v in qc_vals if is_formula(v))
test("100 QC formulas", n_f == 100, f"got {n_f}")

# P8 — formula-aware, focused on PRESERVATION + SOURCE
print("\n[P8] Reply Owner Type — formula preservation")
r44 = find_row(ws, h['Comment ID'], 'IC-0044')
if r44:
    v = ws.cell(row=r44, column=h['Reply Owner Type']).value
    if is_formula(v):
        test("IC-0044 Reply Owner formula preserved",
             'VLOOKUP' in v and 'IFERROR' in v)
        test("formula not hardcoded with 0 literal",
             not any(p.strip() == '0' for p in v.split(',')),
             f"formula length {len(v)}")
    else:
        test("IC-0044 Reply Owner blank", v in (None, ''), f"got {v!r}")

# P10
print("\n[P10] COMPLETENESS_CHECK")
ws_cc = wb['COMPLETENESS_CHECK']; h_cc = hdr(wb, 'COMPLETENESS_CHECK')
for r in range(5, ws_cc.max_row + 1):
    if str(ws_cc.cell(row=r, column=h_cc['Date']).value)[:10] == '2026-07-22':
        st = ws_cc.cell(row=r, column=h_cc['Daily Account Data Status']).value
        test("2026-07-22 = Review", st == 'Review', f"got {st!r}")

# P11
print("\n[P11] EVIDENCE_INDEX")
ws_ei = wb['EVIDENCE_INDEX']; h_ei = hdr(wb, 'EVIDENCE_INDEX')
r = find_row(ws_ei, h_ei['Evidence ID'], 'EVD-SYN-0040-POST')
lr = ws_ei.cell(row=r, column=h_ei['Linked Row']).value if r else None
test("Linked Row = 44", lr == 44, f"got {lr}")

# P12
print("\n[P12] DAILY_BRAND")
ws_db = wb['DAILY_BRAND']; h_db = hdr(wb, 'DAILY_BRAND')
for r in range(5, ws_db.max_row + 1):
    if str(ws_db.cell(row=r, column=h_db['Date']).value)[:10] == '2026-07-24':
        pc = ws_db.cell(row=r, column=h_db['Product Clicks']).value
        test("Product Clicks = 237", pc == 237, f"got {pc}")

# FIX_LOG
print("\n[FIX_LOG]")
test("FIX_LOG exists", 'FIX_LOG' in wb.sheetnames)
if 'FIX_LOG' in wb.sheetnames:
    ws_log = wb['FIX_LOG']
    test("has entries", ws_log.max_row - 1 >= 12, f"got {ws_log.max_row - 1}")
    fix_ids = {str(ws_log.cell(row=r, column=1).value)
               for r in range(2, ws_log.max_row + 1)
               if ws_log.cell(row=r, column=1).value}
    for pid in ['P1', 'P2', 'P3', 'P4', 'P10', 'P11', 'P12']:
        test(f"contains {pid}", pid in fix_ids or f"{pid}fix" in fix_ids)
    test("contains P8 family", any(s.startswith('P8') for s in fix_ids),
         f"ids: {sorted(fix_ids)[:10]}")

# Working copy
print("\n[Working Copy]")
o = wb_orig['DAILY_ACCOUNT'].max_row - 4
w = wb['DAILY_ACCOUNT'].max_row - 4
test("Working != Original", o != w, f"orig={o} work={w}")
test("Original unchanged (39)", o == 39)

print()
print("=" * 70)
print(f"RESULT: {PASS} passed / {FAIL} failed / {PASS + FAIL} total")
print("=" * 70)
sys.exit(0 if FAIL == 0 else 1)
